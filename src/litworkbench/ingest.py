"""PDF / DOCX 导入、检索与阅读。

* 文字层足够的页直接使用提取文字，**不伪造行框**；
* 文字层不足（例如只有页码/页眉的扫描件）时进入 OCR；
* OCR 不可用则单元状态为 ``pending`` 并在 notes 中标注「待识别」；
* 每个单元的 ID 由「文献 ID + 类型 + 顺序」决定，重复导入不会产生副本。
"""

from __future__ import annotations

import re
import shutil
import tempfile
from contextlib import nullcontext
from pathlib import Path
from typing import Any, Iterable

from litworkbench import ocr as ocr_mod
from litworkbench.errors import LitError, NotFoundError, UnsupportedFileError
from litworkbench.store import Workspace, sha256_file, utcnow
from litworkbench.text import (
    LIKE_ESCAPE,
    UnitStatus,
    collapse_pdf_text,
    first_meaningful_chars,
    like_pattern,
    snippet,
    split_terms,
)

SUPPORTED_EXTENSIONS = (".pdf", ".docx")
OCR_MODES = ("auto", "off", "force")
DECISION_NOT_NEEDED = "not_needed"
DECISION_OCR = "ocr"

_PENDING_NOTE = "待识别：该页没有可用文字层，且当前环境未安装 OCR 引擎"
_PENDING_OFF_NOTE = (
    "待识别：文字层不足且以 --ocr off 导入；重新导入或运行 lit import --ocr force 可识别"
)


# --------------------------------------------------------------------- 工具


def _unit_id(doc_id: str, kind: str, seq: int) -> str:
    """单元主键包含完整文献 ID：不同文献绝不共享同一页的 ID。"""
    return f"u_{doc_id}_{kind}_{seq:04d}"


def _doc_id(digest: str) -> str:
    return f"doc_{digest[:16]}"


def _check_extension(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".doc":
        raise UnsupportedFileError(
            "不支持旧版 .doc：请先用 Word 或 LibreOffice 另存为 .docx，再导入。"
        )
    if suffix == ".rtf":
        raise UnsupportedFileError("不支持 .rtf：请先转换为 .docx 或 .pdf。")
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(
            f"不支持的文件类型：{suffix or '（无扩展名）'}；"
            f"当前仅支持 {'、'.join(SUPPORTED_EXTENSIONS)}。"
        )
    return suffix


def decide_ocr(text: str, mode: str, min_chars: int) -> str:
    """逐页决定是否需要 OCR。"""
    if mode == "force":
        return DECISION_OCR
    if mode == "off":
        return DECISION_NOT_NEEDED
    if first_meaningful_chars(text) >= min_chars:
        return DECISION_NOT_NEEDED
    return DECISION_OCR


# ------------------------------------------------------------------- 导入


def import_document(workspace: Workspace, path: Path | str, ocr_mode: str | None = None,
                    title: str | None = None) -> dict[str, Any]:
    """导入单个文件。相同内容哈希会复用已有文献并重试未完成的单元。

    解析过程中**每页/每段立即落盘**：中途失败或被中断时，已经完成的部分
    仍然保留、可诊断、可重试，不会因为一次异常全部丢失。
    """
    source = Path(path).expanduser().resolve()
    if not source.exists():
        raise LitError(f"文件不存在：{source}")
    if not source.is_file():
        raise LitError(f"不是文件：{source}")
    suffix = _check_extension(source)
    if ocr_mode is None:
        ocr_mode = str(workspace.config.get("ocr", {}).get("mode", "auto"))
    if ocr_mode not in OCR_MODES:
        raise LitError(f"未知的 OCR 模式：{ocr_mode}（可选 {'/'.join(OCR_MODES)}）")

    digest = sha256_file(source)
    existing = workspace.document_by_hash(digest)
    if existing:
        _store_copy(workspace, source, digest, suffix)
        retried = retry_pending(workspace, doc_id=existing["id"], ocr_mode=ocr_mode)
        document = workspace.get_document(existing["id"])
        document["duplicate"] = True
        document["retried"] = retried["retried"]
        document["errors"] = retried["errors"]
        if document["status"] in {"failed", "empty"} and retried["errors"]:
            reason = retried["errors"][0].get("error", "未知原因")
            raise LitError(f"重新导入仍然失败（{source.name}）：{reason}")
        return document

    stored = _store_copy(workspace, source, digest, suffix)
    doc_id = _doc_id(digest)
    document = {
        "id": doc_id,
        "sha256": digest,
        "original_path": str(source),
        "stored_path": stored.relative_to(workspace.root).as_posix(),
        "filename": source.name,
        "filetype": suffix.lstrip("."),
        "title": title or "",
        "author": "",
        "year": "",
        "edition": "",
        "volume": "",
        "status": "importing",
        "ocr_mode": ocr_mode,
        "page_count": 0,
        "unit_count": 0,
        "pending_count": 0,
    }
    workspace.save_document(document)
    return _extract_into(workspace, document, stored, suffix, ocr_mode)


def _extract_into(workspace: Workspace, document: dict[str, Any], stored: Path,
                  suffix: str, ocr_mode: str,
                  skip_completed: set[str] | None = None) -> dict[str, Any]:
    """逐页/逐段解析并立即保存；异常时保留已完成部分并给出真实状态。

    ``skip_completed`` 中的单元 ID 会被跳过（用于补做缺失页时不重做已完成页）。
    """
    doc_id = document["id"]
    workspace.update_document(doc_id, status="importing", ocr_mode=ocr_mode)
    counts = 0
    try:
        if suffix == ".pdf":
            blocks = _iter_pdf_blocks(workspace, stored, doc_id, ocr_mode, skip_completed)
        else:
            blocks = _iter_docx_blocks(workspace, stored, doc_id, ocr_mode)
        metadata: dict[str, Any] = {}
        for block, block_metadata in blocks:
            metadata.update({k: v for k, v in (block_metadata or {}).items() if v})
            if not block.get("id"):
                continue
            _save_extraction(workspace, block)
            counts += 1
            if metadata:
                # 页数等信息边做边存，中断后仍然知道文档原本有多少页
                workspace.update_document(doc_id, **metadata)
    except KeyboardInterrupt:
        workspace.update_document(doc_id, **_final_status(workspace, doc_id, "interrupted"))
        raise
    except LitError:
        workspace.update_document(doc_id, **_final_status(workspace, doc_id, "failed"))
        raise
    except Exception as exc:  # pragma: no cover - 损坏文件等
        workspace.update_document(doc_id, **_final_status(workspace, doc_id, "failed"))
        raise LitError(f"解析失败（{document['filename']}）：{type(exc).__name__}: {exc}") from exc

    updates: dict[str, Any] = dict(metadata)
    if not document.get("title"):
        updates["title"] = Path(document["filename"]).stem
    updates.update(_final_status(workspace, doc_id, ""))
    return workspace.update_document(doc_id, **updates)


def _final_status(workspace: Workspace, doc_id: str, forced: str, *, connection=None) -> dict[str, Any]:
    """真实状态：零单元、有未完成单元或被中断的文献都**不会**是 completed。"""
    units = workspace.units_for_doc(doc_id, connection=connection)
    pending = sum(1 for unit in units if unit["status"] != UnitStatus.COMPLETED)
    document = workspace.get_document(doc_id, connection=connection)
    expected = int(document.get("page_count") or 0) if document["filetype"] == "pdf" else len(units)
    missing = max(0, expected - len(units))
    pending += missing
    if forced:
        status = forced
    elif not units:
        status = "empty"
    elif pending == 0:
        status = "completed"
    else:
        status = "partial"
    return {"status": status, "unit_count": len(units), "pending_count": pending}


def _store_copy(workspace: Workspace, source: Path, digest: str, suffix: str) -> Path:
    target_dir = workspace.sources_dir / digest[:2]
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{digest}{suffix}"
    if target.exists() and sha256_file(target) == digest:
        return target
    with tempfile.NamedTemporaryFile(dir=target_dir, suffix=".part", delete=False) as handle:
        staging = Path(handle.name)
    try:
        shutil.copy2(source, staging)
        if sha256_file(staging) != digest:
            raise LitError("复制期间源文件发生变化，请重新导入。")
        staging.replace(target)
    finally:
        staging.unlink(missing_ok=True)
    return target


def _save_extraction(workspace: Workspace, fresh: dict[str, Any], *, connection=None) -> None:
    """保存提取结果，并在同一事务内保留旧版本和更新证据核验状态。"""
    from litworkbench.evidence import _mark_stale_evidence

    with (nullcontext(connection) if connection is not None else workspace.connect()) as conn:
        if connection is None:
            conn.execute("BEGIN IMMEDIATE")
        fresh["extraction_status"] = fresh["status"]
        previous = conn.execute("SELECT * FROM units WHERE id = ?", (fresh["id"],)).fetchone()
        if previous:
            fresh["corrected_text"] = previous["corrected_text"]
            fresh["printed_page"] = previous["printed_page"]
            fresh["created_at"] = previous["created_at"]
            if fresh["corrected_text"] is not None:
                fresh["status"] = "completed" if fresh["corrected_text"].strip() else "pending"
        workspace.save_unit(fresh, connection=conn)
        if previous and previous["original_text"] != fresh["original_text"]:
            workspace.add_revision(fresh["id"], "extraction", previous["original_text"],
                                   fresh["original_text"], "重新提取或补做 OCR", connection=conn)
            # 按每条摘录**记录时的版本**判断是否仍成立（原文变了才让原文摘录需复核）。
            _mark_stale_evidence(workspace, fresh["id"], connection=conn)


def _iter_pdf_blocks(workspace: Workspace, path: Path, doc_id: str, ocr_mode: str,
                     skip_completed: set[str] | None = None):
    """逐页产出 ``(block, metadata)``；调用方负责立刻落盘。

    CLI 与后台队列共用 :func:`open_pdf_pages`，因此「命令行导入」和
    「网页后台识别」的逐页判定完全一致。
    """
    with open_pdf_pages(workspace, path, doc_id, ocr_mode) as pages:
        metadata = dict(pages.metadata)
        if pages.page_count == 0:
            raise LitError(f"PDF 无法解析：{path.name}（文件可能已损坏或被加密）")
        for index in range(pages.page_count):
            unit_id = _unit_id(doc_id, "page", index + 1)
            if skip_completed and unit_id in skip_completed:
                yield {}, {"page_count": pages.page_count}
                continue
            yield pages.block_for(index), {"page_count": pages.page_count}
        yield {}, metadata


class PdfPages:
    """打开一个 PDF，按页提供待保存的 block；关闭后释放 PDFium 句柄。

    后台队列逐页提交检查点时使用它，从而做到「每页结束即落盘」，
    而不是把整本书的 OCR 攒在一次调用里。
    """

    def __init__(self, workspace: Workspace, path: Path, doc_id: str, ocr_mode: str,
                 *, ocr_page=None) -> None:
        import pypdfium2 as pdfium

        self._workspace = workspace
        self._path = Path(path)
        self._doc_id = doc_id
        self._ocr_mode = ocr_mode
        self._ocr_page = ocr_page or _ocr_page
        self._config = workspace.config.get("ocr", {})
        self._scale = float(self._config.get("render_scale", 2.0))
        self._min_chars = int(self._config.get("min_text_chars", 20))
        self._pdf = pdfium.PdfDocument(str(self._path))
        self.page_count = len(self._pdf)
        self.metadata = self._read_metadata()

    @property
    def scale(self) -> float:
        return self._scale

    @property
    def min_chars(self) -> int:
        return self._min_chars

    def _read_metadata(self) -> dict[str, Any]:
        try:
            raw = self._pdf.get_metadata_dict() or {}
            return {
                "title": str(raw.get("Title") or "").strip(),
                "author": str(raw.get("Author") or "").strip(),
            }
        except Exception:  # pragma: no cover - 元数据可有可无
            return {}

    def page_text(self, index: int) -> str:
        page = self._pdf[index]
        try:
            textpage = page.get_textpage()
            try:
                return collapse_pdf_text(textpage.get_text_range() or "")
            finally:
                textpage.close()
        finally:
            page.close()

    def block_for(self, index: int, *, ocr_mode: str | None = None) -> dict[str, Any]:
        """构造第 ``index`` 页（0 基）的 block；不落盘，由调用方决定提交时机。"""
        mode = ocr_mode or self._ocr_mode
        text = self.page_text(index)
        seq = index + 1
        decision = decide_ocr(text, mode, self._min_chars)
        sparse = first_meaningful_chars(text) < self._min_chars
        block: dict[str, Any] = {
            "id": _unit_id(self._doc_id, "page", seq),
            "doc_id": self._doc_id,
            "seq": seq,
            "kind": "page",
            "file_page": seq,
            "printed_page": "",
            "title_path": "",
            "method": "text",
            "status": UnitStatus.COMPLETED,
            "ocr_decision": decision,
            "original_text": text,
            "corrected_text": None,
            "notes": "",
            "ocr_json": None,
        }
        if decision == DECISION_OCR:
            block.update(self._ocr_page(self._workspace, self._path, index, text, self._scale))
        elif sparse:
            # OCR 被显式关闭：仍然是「待识别」，不冒充已完成
            block["status"] = UnitStatus.PENDING
            block["notes"] = _PENDING_OFF_NOTE
        return block

    def close(self) -> None:
        try:
            self._pdf.close()
        except Exception:  # pragma: no cover - 关闭失败不影响已提交的结果
            pass

    def __enter__(self) -> "PdfPages":
        return self

    def __exit__(self, *exc_info: Any) -> None:
        self.close()


def open_pdf_pages(workspace: Workspace, path: Path | str, doc_id: str,
                   ocr_mode: str, *, ocr_page=None) -> PdfPages:
    return PdfPages(workspace, Path(path), doc_id, ocr_mode, ocr_page=ocr_page)


def failed_page_block(doc_id: str, file_page: int, reason: str,
                      status: str = UnitStatus.FAILED) -> dict[str, Any]:
    """构造一个「这一页没拿到可用文字」的占位 block。

    页面提取抛异常时用它落盘，这样文献里**真的存在**这一条失败单元：
    覆盖、任务计数与重试都以数据库为准，不会因为缺行就把失败页算成成功。
    """
    seq = int(file_page)
    return {
        "id": _unit_id(doc_id, "page", seq),
        "doc_id": doc_id,
        "seq": seq,
        "kind": "page",
        "file_page": seq,
        "printed_page": "",
        "title_path": "",
        "method": "text",
        "status": status,
        "ocr_decision": DECISION_OCR,
        "original_text": "",
        "corrected_text": None,
        "notes": reason,
        "ocr_json": None,
    }


def force_page_block(workspace: Workspace, path: Path | str, doc_id: str,
                     file_page: int, ocr_mode: str = "force", *, ocr_page=None) -> dict[str, Any]:
    """对已存在的某页**强制**重新提取/OCR，返回待保存的 block。

    与普通重试的区别：已经 completed 的页也会真正执行 OCR。
    调用方负责在新版本提交前保留旧结果。
    """
    index = int(file_page) - 1
    with open_pdf_pages(workspace, path, doc_id, "off") as pages:
        text = pages.page_text(index)
        seq = index + 1
        block: dict[str, Any] = {
            "id": _unit_id(doc_id, "page", seq),
            "doc_id": doc_id,
            "seq": seq,
            "kind": "page",
            "file_page": seq,
            "printed_page": "",
            "title_path": "",
            "method": "text",
            "status": UnitStatus.COMPLETED,
            "ocr_decision": DECISION_OCR,
            "original_text": text,
            "corrected_text": None,
            "notes": "",
            "ocr_json": None,
        }
        block.update((ocr_page or _ocr_page)(workspace, Path(path), index, text, pages.scale))
        return block


def _ocr_page(workspace: Workspace, pdf_path: Path, index: int, text: str,
              scale: float) -> dict[str, Any]:
    """对单页执行真实 OCR；引擎不可用时返回 pending 状态。"""
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(pdf_path))
        try:
            page = pdf[index]
            try:
                bitmap = page.render(scale=scale)
                image = bitmap.to_pil().convert("RGB")
                try:
                    result = ocr_mod.recognize(image)
                finally:
                    image.close()
            finally:
                page.close()
        finally:
            pdf.close()
    except ocr_mod.OcrUnavailable:
        return {
            "method": "text",
            "status": UnitStatus.PENDING,
            "notes": _PENDING_NOTE,
            "original_text": text,
        }
    except Exception as exc:  # pragma: no cover - 渲染异常
        return {
            "method": "text",
            "status": UnitStatus.FAILED,
            "notes": f"OCR 失败：{type(exc).__name__}: {exc}",
            "original_text": text,
        }

    notes = ""
    if result["mean_confidence"] < 0.80:
        notes = f"OCR 平均置信度偏低（{result['mean_confidence']}），建议人工核对"
    if not result["text"].strip():
        return {
            "method": "ocr",
            "status": UnitStatus.PENDING,
            "notes": "待识别：OCR 未返回任何文字",
            "original_text": text,
            "ocr_json": result,
        }
    return {
        "method": "ocr",
        "status": UnitStatus.COMPLETED,
        "notes": notes,
        "original_text": result["text"],
        "ocr_json": result,
    }


def _iter_docx_blocks(workspace: Workspace, path: Path, doc_id: str, ocr_mode: str):
    """按文档顺序产出段落与表格单元（DOCX 不编造页码，也不需要 OCR）。"""
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    core = document.core_properties
    metadata = {
        "title": (core.title or "").strip(),
        "author": (core.author or "").strip(),
        "year": "",  # 电子文件创建时间不能作为书籍出版年份。
    }

    seq = 0
    heading_stack: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _heading_level(paragraph)
            if level:
                heading_stack = heading_stack[: level - 1]
                heading_stack.append(text)
            seq += 1
            # title_path = 该单元的祖先标题（不包含它自己）
            trail = heading_stack[:-1] if level else heading_stack
            yield {
                "id": _unit_id(doc_id, "para", seq),
                "doc_id": doc_id,
                "seq": seq,
                "kind": "paragraph",
                "file_page": None,
                "printed_page": "",
                "title_path": " / ".join(trail),
                "method": "text",
                "status": UnitStatus.COMPLETED,
                "ocr_decision": DECISION_NOT_NEEDED,
                "original_text": text,
                "corrected_text": None,
                "notes": "",
                "ocr_json": None,
            }, None
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    rows.append("\t".join(cells))
            if not rows:
                continue
            seq += 1
            yield {
                "id": _unit_id(doc_id, "table", seq),
                "doc_id": doc_id,
                "seq": seq,
                "kind": "table",
                "file_page": None,
                "printed_page": "",
                "title_path": " / ".join(heading_stack),
                "method": "text",
                "status": UnitStatus.COMPLETED,
                "ocr_decision": DECISION_NOT_NEEDED,
                "original_text": "\n".join(rows),
                "corrected_text": None,
                "notes": "",
                "ocr_json": None,
            }, None
    yield {}, metadata


def _heading_level(paragraph: Any) -> int:
    style = getattr(paragraph, "style", None)
    name = getattr(style, "name", "") or ""
    match = re.search(r"(\d+)", name)
    if match and ("Heading" in name or "标题" in name):
        return int(match.group(1))
    return 0


def import_path(workspace: Workspace, path: Path | str, ocr_mode: str | None = None,
                recursive: bool = True) -> dict[str, Any]:
    """导入文件或目录；目录只接受允许的扩展名，且不会递归进工作区自身。"""
    target = Path(path).expanduser().resolve()
    if not target.exists():
        raise LitError(f"路径不存在：{target}")
    results: list[dict[str, Any]] = []
    skipped: list[str] = []
    errors: list[dict[str, str]] = []

    if target.is_file():
        candidates: Iterable[Path] = [target]
    else:
        pattern = "**/*" if recursive else "*"
        candidates = [item for item in sorted(target.glob(pattern)) if item.is_file()]

    workspace_root = workspace.root.resolve()
    for candidate in candidates:
        if not target.is_file() and (_is_inside(candidate, workspace_root) or candidate.suffix.lower() not in SUPPORTED_EXTENSIONS):
            skipped.append(str(candidate))
            continue
        try:
            results.append(import_document(workspace, candidate, ocr_mode=ocr_mode))
        except LitError as exc:
            errors.append({"path": str(candidate), "error": str(exc)})

    return {
        "root": str(target),
        "imported": len(results),
        "skipped": len(skipped),
        "errors": errors,
        "documents": results,
    }


def _is_inside(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent)
    except ValueError:
        return False
    return True


def retry_pending(workspace: Workspace, doc_id: str | None = None,
                  ocr_mode: str | None = None) -> dict[str, Any]:
    """重新处理未完成的工作。

    * 解析失败/中断且**没有留下任何单元**的文献：整体重新解析；
    * 已完成单元不受影响，只补做 ``pending`` / ``failed`` 单元；
    * 始终按真实情况写回状态，失败不会变成 completed。
    """
    if ocr_mode is None:
        ocr_mode = str(workspace.config.get("ocr", {}).get("mode", "auto"))
    docs = [workspace.get_document(doc_id)] if doc_id else workspace.list_documents(limit=1000)
    retried = 0
    reextracted: list[str] = []
    failed: list[dict[str, str]] = []
    for document in docs:
        if document["filetype"] != "pdf":
            continue
        path = workspace.path_of(document)
        units = workspace.units_for_doc(document["id"])
        completed = {unit["id"] for unit in units if unit["status"] == UnitStatus.COMPLETED}
        page_count = int(document.get("page_count") or 0)
        missing_pages = bool(units) and page_count > len(units)

        if not units:
            try:
                _extract_into(workspace, document, path, ".pdf", ocr_mode)
                reextracted.append(document["id"])
                retried += len(workspace.units_for_doc(document["id"]))
            except LitError as exc:
                workspace.update_document(document["id"], status="failed")
                failed.append({"doc_id": document["id"], "error": str(exc)})
            except Exception as exc:  # pragma: no cover - 保护性分支
                workspace.update_document(document["id"], status="failed")
                failed.append({"doc_id": document["id"], "error": f"{type(exc).__name__}: {exc}"})
            continue

        if missing_pages:
            # 上次被中断：只补做缺失的页，已完成的页不重做
            try:
                _extract_into(workspace, document, path, ".pdf", ocr_mode,
                              skip_completed=completed)
                reextracted.append(document["id"])
                retried += len(workspace.units_for_doc(document["id"])) - len(units)
            except LitError as exc:
                failed.append({"doc_id": document["id"], "error": str(exc)})
            except Exception as exc:  # pragma: no cover - 保护性分支
                failed.append({"doc_id": document["id"], "error": f"{type(exc).__name__}: {exc}"})
            units = workspace.units_for_doc(document["id"])

        for unit in units:
            if unit["status"] == UnitStatus.COMPLETED:
                continue
            try:
                fresh = _reprocess_pdf_unit(workspace, path, document, unit, ocr_mode)
                _save_extraction(workspace, fresh)
                retried += 1
            except Exception as exc:  # pragma: no cover - 保护性分支
                workspace.mark_unit_failed(unit["id"], f"重试失败：{type(exc).__name__}: {exc}")
                failed.append({"unit_id": unit["id"], "error": str(exc)})
        workspace.update_document(
            document["id"], **_final_status(workspace, document["id"], "")
        )
    return {
        "retried": retried,
        "reextracted": reextracted,
        "documents": [doc["id"] for doc in docs],
        "errors": failed,
    }


def _reprocess_pdf_unit(workspace: Workspace, path: Path, document: dict[str, Any],
                        unit: dict[str, Any], ocr_mode: str) -> dict[str, Any]:
    import pypdfium2 as pdfium

    index = int(unit["file_page"]) - 1
    pdf = pdfium.PdfDocument(str(path))
    try:
        page = pdf[index]
        try:
            textpage = page.get_textpage()
            try:
                raw = textpage.get_text_range() or ""
            finally:
                textpage.close()
            text = collapse_pdf_text(raw)
            cfg_ocr = workspace.config.get("ocr", {})
            decision = decide_ocr(
                text, ocr_mode, int(cfg_ocr.get("min_text_chars", 20))
            )
            fresh = dict(unit)
            fresh["notes"] = ""
            fresh["corrected_text"] = unit.get("corrected_text")
            if decision == DECISION_OCR:
                scale = float(cfg_ocr.get("render_scale", 2.0))
                fresh.update(_ocr_page(workspace, path, index, text, scale))
                fresh["ocr_decision"] = DECISION_OCR
            else:
                fresh["original_text"] = text
                fresh["method"] = "text"
                fresh["ocr_json"] = None
                fresh["ocr_decision"] = decision
                if first_meaningful_chars(text) < int(cfg_ocr.get("min_text_chars", 20)):
                    fresh["status"] = UnitStatus.PENDING
                    fresh["notes"] = _PENDING_OFF_NOTE
                else:
                    fresh["status"] = UnitStatus.COMPLETED
            return fresh
        finally:
            page.close()
    finally:
        pdf.close()


# ------------------------------------------------------------------- 检索


def search(workspace: Workspace, query: str, limit: int = 20,
           include_pending: bool = False) -> list[dict[str, Any]]:
    """本地中文文字检索；空白与换行不参与匹配，``%``/``_`` 按字面处理。"""
    terms = split_terms(query)
    if not terms:
        raise LitError("检索词为空：请输入至少一个非空白字符。")
    limit = max(1, int(limit))
    statuses = ("completed",)
    if include_pending:
        statuses = UnitStatus.ALL
    where = [f"u.status IN ({', '.join('?' * len(statuses))})"]
    params: list[Any] = list(statuses)
    for term in terms:
        pattern = like_pattern(term)
        where.append(
            f"(CASE WHEN u.corrected_text IS NOT NULL THEN u.search_text_corrected "
            f"ELSE u.search_text END) LIKE ? ESCAPE '{LIKE_ESCAPE}'"
        )
        params.append(pattern)
    sql = (
        "SELECT u.*, d.filename, d.title AS doc_title, d.filetype "
        "FROM units u JOIN documents d ON d.id = u.doc_id "
        f"WHERE {' AND '.join(where)} ORDER BY d.id, u.seq LIMIT ?"
    )
    params.append(limit * 8)
    with workspace.connect() as conn:
        rows = conn.execute(sql, params).fetchall()

    hits: list[dict[str, Any]] = []
    for row in rows:
        unit = dict(row)
        unit.pop("search_text", None)
        unit.pop("search_text_corrected", None)
        effective, version = _effective_text(unit)
        matched_term = next((t for t in terms if snippet(effective, t, 1)), terms[0])
        hits.append(
            {
                "unit_id": unit["id"],
                "doc_id": unit["doc_id"],
                "filename": unit["filename"],
                "title": unit["doc_title"],
                "kind": unit["kind"],
                "seq": unit["seq"],
                "file_page": unit["file_page"],
                "printed_page": unit["printed_page"],
                "method": unit["method"],
                "status": unit["status"],
                "version": version,
                "notes": unit["notes"],
                "term": matched_term,
                "snippet": snippet(effective, matched_term),
            }
        )
        if len(hits) >= limit:
            break
    return hits


def _effective_text(unit: dict[str, Any]) -> tuple[str, str]:
    if unit.get("corrected_text") is not None:
        return unit["corrected_text"], "corrected"
    return unit.get("original_text") or "", "original"


def read_unit(workspace: Workspace, unit_id: str) -> dict[str, Any]:
    """读取一个原文单元，返回定位信息与文本。"""
    unit = workspace.get_unit(unit_id)
    document = workspace.get_document(unit["doc_id"])
    text, version = _effective_text(unit)
    return {
        "unit_id": unit["id"],
        "doc_id": unit["doc_id"],
        "filename": document["filename"],
        "title": document["title"],
        "kind": unit["kind"],
        "seq": unit["seq"],
        "file_page": unit["file_page"],
        "printed_page": unit["printed_page"],
        "title_path": unit["title_path"],
        "method": unit["method"],
        "status": unit["status"],
        "notes": unit["notes"],
        "version": version,
        "text": text,
        "original_text": unit["original_text"],
        "corrected_text": unit.get("corrected_text"),
        "has_correction": unit.get("corrected_text") is not None,
        "ocr": unit.get("ocr_json") or None,
        "revisions": [
            {
                "at": revision["at"],
                "kind": revision["kind"],
                "reason": revision["reason"],
            }
            for revision in workspace.revisions_for_unit(unit["id"])
        ],
    }


def render_page(workspace: Workspace, unit_id: str, scale: float = 2.0,
                output: Path | str | None = None) -> dict[str, Any]:
    """把 PDF 页渲染成图片（默认写入 artifacts/pages），便于人工核验。"""
    unit = workspace.get_unit(unit_id)
    if unit["kind"] != "page":
        raise LitError(f"单元 {unit_id} 不是 PDF 页面，无法渲染页图。")
    document = workspace.get_document(unit["doc_id"])
    if document["filetype"] != "pdf":
        raise LitError(f"文献 {document['id']} 不是 PDF。")
    import pypdfium2 as pdfium

    if output is None:
        workspace.renders_dir.mkdir(parents=True, exist_ok=True)
        target = workspace.renders_dir / f"{unit_id}_s{scale:g}.png"
    else:
        target = Path(output).expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)

    pdf = pdfium.PdfDocument(str(workspace.path_of(document)))
    try:
        page = pdf[int(unit["file_page"]) - 1]
        try:
            image = page.render(scale=float(scale)).to_pil().convert("RGB")
            try:
                image.save(target)
                size = image.size
            finally:
                image.close()
        finally:
            page.close()
    finally:
        pdf.close()
    return {
        "unit_id": unit_id,
        "path": target,
        "width": size[0],
        "height": size[1],
        "scale": float(scale),
        "rendered_at": utcnow(),
    }


def list_documents(workspace: Workspace, limit: int = 100) -> list[dict[str, Any]]:
    return workspace.list_documents(limit=limit)


def read_document(workspace: Workspace, needle: str) -> dict[str, Any]:
    document = workspace.find_document(needle)
    document["units"] = workspace.units_for_doc(document["id"])
    return document


def require_unit(workspace: Workspace, unit_id: str) -> dict[str, Any]:
    try:
        return workspace.get_unit(unit_id)
    except NotFoundError:
        raise
