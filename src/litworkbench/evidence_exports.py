"""网页资料导出：一次一致快照 → 四种可下载的资料集文件。

契约要点：

* **范围**：未指定核验状态时**只导出 accepted**；显式 ``all`` 才包含全部状态，
  并且逐条保留核验状态与「复核前不应引用」的提示；
* **显式选择与筛选取交集**：未知 ID 直接拒绝，不静默忽略；
* **内容摘要（snapshot_token）** 覆盖格式、完整范围、每条摘录的字段、书目与页码，
  但**不包含生成时间**；预览与生成之间只要内容变了，生成就返回 409，
  不会把上一次预览时的核验判断当成本次结果；
* **历史引文不被改写**：引文、偏移与文字哈希取自摘录当时保存的版本快照，
  不用当前全文替换；书目与当前核验状态来自同一个读事务；
* **不泄露本机绝对路径**：文件里没有原件的本机路径，只有文件名与内容哈希；
* **文件由服务端决定**：文件名来自摘要，路径固定在 ``artifacts/exports``，
  下载只按数据库里的 ID 解析，没有任意路径参数。

这四种文件都是**资料集**，不是研究报告：不生成结论、不评价文献、不做归纳。
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import sqlite3
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from litworkbench.collections import (
    assert_collection_exists,
    collection_ids_for,
    normalize_ids,
)
from litworkbench.errors import EditConflictError, NotFoundError, ValidationError
from litworkbench.evidence import (
    REVIEW_ACCEPTED,
    REVIEW_REJECTED,
    REVIEW_STALE,
    REVIEW_STATUSES,
)
from litworkbench.export import evidence_repository, evidence_view, csv_safe, md_escape
from litworkbench.store import Workspace, sha256_file, sha256_text, utcnow

FORMATS = ("csv", "json", "markdown", "docx")
EXTENSIONS = {"csv": ".csv", "json": ".json", "markdown": ".md", "docx": ".docx"}
MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "json": "application/json; charset=utf-8",
    "markdown": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
FORMAT_LABELS = {"csv": "CSV 表格", "json": "JSON 数据", "markdown": "Markdown 文本",
                 "docx": "Word 文档"}

#: 显式 ``all`` 表示包含全部核验状态
SCOPE_ALL = "all"
DEFAULT_REVIEW = REVIEW_ACCEPTED
#: 读取时逐批取，避免一次把整库塞进内存；**不会**因此截断结果
BATCH_SIZE = 500

KIND = "excerpt_material_set"

MATERIAL_NOTE = (
    "本文件是摘录资料集：每条都是原文的精确切片，附来源、核验状态与核验记录。"
    "它只整理已保存的摘录，不包含研究结论、归纳或评价。"
)
REVIEW_NOTE = (
    "「已接受」只表示该条摘录已被人工对照原页核验，并不表示任何研究结论已被确认；"
    "标注为需复核（stale）或已拒绝（rejected）的条目在人工复核之前不应引用。"
)

STATUS_LABELS = {
    "pending": "待核验",
    REVIEW_ACCEPTED: "已接受",
    REVIEW_REJECTED: "已拒绝",
    REVIEW_STALE: "需复核",
}

#: 每条摘录在四种格式里都出现的字段（CSV 表头即此顺序）
MATERIAL_COLUMNS = (
    "evidence_id", "task_id", "doc_id", "unit_id", "unit_seq", "kind",
    "filename", "title", "author", "year", "edition", "volume",
    "source_sha256", "file_page", "printed_page", "printed_page_status",
    "title_path", "position", "locator", "version", "unit_text_sha256", "method",
    "unit_text_length", "recorded_at", "start_offset", "end_offset", "quote",
    "category", "note", "review_status", "review_status_label", "review_note",
    "collections", "usable", "warning", "created_at",
)


# ------------------------------------------------------------------ 范围


@dataclass(frozen=True)
class Scope:
    """一次导出的完整范围。解析后不可变，预览与生成共用同一份语义。"""

    doc_id: str | None = None
    category: str | None = None
    review_status: str | None = None
    collection_id: str | None = None
    evidence_ids: tuple[str, ...] | None = None

    @property
    def review_policy(self) -> str:
        """``accepted``（归一化后的默认）或 ``all`` / 具体状态。"""
        return self.review_status or DEFAULT_REVIEW

    @property
    def includes_all_statuses(self) -> bool:
        return self.review_status == SCOPE_ALL

    def as_dict(self) -> dict[str, Any]:
        return {
            "doc_id": self.doc_id,
            "category": self.category,
            "review_status": self.review_status,
            "collection_id": self.collection_id,
            "evidence_ids": list(self.evidence_ids) if self.evidence_ids is not None else None,
        }


def normalize_scope(*, doc_id: str | None = None, category: str | None = None,
                    review_status: str | None = None, collection_id: str | None = None,
                    evidence_ids: list[str] | None = None,
                    fmt: str | None = None) -> Scope:
    """校验范围。未知值一律在这里变成中文错误，不进入任何查询。

    核验口径在**这里**定死：不给 ``review_status`` 就归一化成 ``accepted``，
    只有显式传 ``all``（或某个具体状态）才会绕过默认值。
    因此不存在「范围说只含已接受、查询却把所有状态都导出来」的空档。
    """
    if fmt is not None and fmt not in FORMATS:
        raise ValidationError(f"不支持的导出格式：{fmt}（可选 {'/'.join(FORMATS)}）")
    status = (review_status or "").strip() or DEFAULT_REVIEW
    if status != SCOPE_ALL and status not in REVIEW_STATUSES:
        raise ValidationError(
            f"未知的核验状态：{review_status}"
            f"（可选 {SCOPE_ALL} 或 {'/'.join(REVIEW_STATUSES)}）"
        )
    # 显式 [] 必须报错，绝不回退成「全部结果」
    ids = (None if evidence_ids is None
           else tuple(normalize_ids(evidence_ids)))
    return Scope(
        doc_id=(doc_id or "").strip() or None,
        category=(category or "").strip() or None,
        review_status=status,
        collection_id=(collection_id or "").strip() or None,
        evidence_ids=ids,
    )


# ------------------------------------------------------------------ 快照


@dataclass(frozen=True)
class Snapshot:
    """一次导出要写进文件的确切内容（与生成时间无关）。

    **完全脱离数据库**：证据、书目、单元、页码与集合名称都在
    :func:`build_snapshot` 的那一个读事务里读进 Python 结构。
    之后的摘要计算、四种序列化与落盘都只读这个对象，
    不会重新查询，因此「算完摘要再改书目/校订原文」也不会出现新旧混排。
    """

    format: str
    scope: Scope
    scope_description: str
    count: int
    review_counts: dict[str, int]
    items: tuple[dict[str, Any], ...]
    documents: tuple[dict[str, Any], ...]
    token: str
    filename: str


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _scope_description(conn: sqlite3.Connection, scope: Scope) -> str:
    parts: list[str] = []
    if scope.doc_id:
        row = conn.execute("SELECT filename, title FROM documents WHERE id = ?",
                           (scope.doc_id,)).fetchone()
        label = (row["title"] or row["filename"]) if row else scope.doc_id
        parts.append(f"来源文献：{label}")
    if scope.collection_id:
        row = conn.execute("SELECT name FROM collections WHERE id = ?",
                           (scope.collection_id,)).fetchone()
        parts.append(f"集合：{row['name'] if row else scope.collection_id}")
    if scope.category:
        parts.append(f"分类：{scope.category}")
    if scope.evidence_ids:
        parts.append(f"显式选择：{len(scope.evidence_ids)} 条摘录（与上述条件取交集）")
    else:
        parts.append("范围：当前筛选的全部结果")
    if scope.includes_all_statuses:
        parts.append("核验状态：全部（逐条保留当前状态，未核验条目不得直接引用）")
    elif scope.review_status == DEFAULT_REVIEW:
        parts.append("核验状态：仅已接受（默认）")
    else:
        parts.append(f"核验状态：{STATUS_LABELS.get(scope.review_status, scope.review_status)}")
    return "；".join(parts)


def select_rows(workspace: Workspace, scope: Scope,
                connection: sqlite3.Connection) -> list[dict[str, Any]]:
    """按范围取全部证据行（分批读取，不受列表页 30 条限制）。"""
    clauses: list[str] = []
    params: list[Any] = []
    if scope.evidence_ids is not None:
        placeholders = ", ".join("?" * len(scope.evidence_ids))
        clauses.append(f"id IN ({placeholders})")
        params.extend(scope.evidence_ids)
        known = {
            row["id"] for row in connection.execute(
                f"SELECT id FROM evidence WHERE id IN ({placeholders})",
                list(scope.evidence_ids),
            ).fetchall()
        }
        unknown = [value for value in scope.evidence_ids if value not in known]
        if unknown:
            raise ValidationError(
                "找不到这些摘录：" + "、".join(unknown[:10])
                + ("…" if len(unknown) > 10 else "")
                + "；请重新选择后再导出。"
            )
    if scope.doc_id:
        clauses.append("doc_id = ?")
        params.append(scope.doc_id)
    if scope.category:
        clauses.append("category = ?")
        params.append(scope.category)
    if scope.review_status and not scope.includes_all_statuses:
        clauses.append("review_status = ?")
        params.append(scope.review_status)
    if scope.collection_id:
        clauses.append("id IN (SELECT evidence_id FROM collection_evidence "
                       "WHERE collection_id = ?)")
        params.append(scope.collection_id)
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    sql = f"SELECT * FROM evidence {where} ORDER BY created_at, id LIMIT ? OFFSET ?"
    rows: list[dict[str, Any]] = []
    offset = 0
    while True:
        batch = connection.execute(sql, [*params, BATCH_SIZE, offset]).fetchall()
        rows.extend(dict(row) for row in batch)
        if len(batch) < BATCH_SIZE:
            break
        offset += BATCH_SIZE
    return rows


def _material_item(view: dict[str, Any], collection_names: list[str]) -> dict[str, Any]:
    """一条摘录的资料行：只含权威字段，不含本机路径。

    「记录时」的字段一律取自 :attr:`version_snapshot`——摘录当时用的文字版本、
    那段文字的哈希和识别方式。只有极老的记录没有快照时才退回单元当前值，
    绝不用**当前**单元的 method 去改写历史摘录的识别方式。
    """
    snapshot = view["version_snapshot"] or {}
    status = view["review_status"]
    usable = bool(view["usable"])
    return {
        "evidence_id": view["id"],
        "task_id": view.get("task_id") or "",
        "doc_id": view["doc_id"],
        "unit_id": view["unit_id"],
        "unit_seq": view["unit_seq"],
        "kind": view["kind"],
        "filename": view["filename"],
        "title": view["title"],
        "author": view["author"],
        "year": view["year"],
        "edition": view["edition"],
        "volume": view["volume"],
        "source_sha256": view["source_sha256"],
        "file_page": view["file_page"],
        "printed_page": view["printed_page"],
        "printed_page_status": view["printed_page_status"],
        "title_path": view["title_path"],
        "position": view["position"],
        "locator": view["locator"],
        "version": view["version"],
        # 历史哈希与历史识别方式：来自摘录当时保存的版本快照，不是当前单元
        "unit_text_sha256": snapshot.get("unit_text_sha256", ""),
        "method": snapshot.get("method") or view["method"],
        "unit_text_length": snapshot.get("unit_text_length"),
        "recorded_at": snapshot.get("captured_at", ""),
        "start_offset": view["start_offset"],
        "end_offset": view["end_offset"],
        "quote": view["quote"],
        "category": view["category"],
        "note": view["note"],
        "review_status": status,
        "review_status_label": STATUS_LABELS.get(status, status),
        "review_note": view.get("review_note", ""),
        "collections": "、".join(collection_names),
        "usable": usable,
        "warning": "" if usable else _warning_text(status),
        "created_at": view["created_at"],
    }


def _warning_text(status: str) -> str:
    if status == REVIEW_STALE:
        return "原文在校订或重新识别后已变化，该引文仍指向记录时的版本，复核前不应直接引用。"
    if status == REVIEW_REJECTED:
        return "该摘录已被人工拒绝，不应引用。"
    if status == "pending":
        return "该摘录尚未人工核验，引用前请先对照原页。"
    return "该摘录当前状态为 %s。" % status


def _document_summary(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """按文献汇总书目（同一快照内读到的值）。"""
    order: list[str] = []
    grouped: dict[str, dict[str, Any]] = {}
    for item in items:
        doc_id = item["doc_id"]
        if doc_id not in grouped:
            order.append(doc_id)
            grouped[doc_id] = {
                "doc_id": doc_id,
                "filename": item["filename"],
                "title": item["title"],
                "author": item["author"],
                "year": item["year"],
                "edition": item["edition"],
                "volume": item["volume"],
                "source_sha256": item["source_sha256"],
                "count": 0,
            }
        grouped[doc_id]["count"] += 1
    return [grouped[doc_id] for doc_id in order]


def _review_counts(items: list[dict[str, Any]]) -> dict[str, int]:
    counts = {status: 0 for status in REVIEW_STATUSES}
    for item in items:
        counts[item["review_status"]] = counts.get(item["review_status"], 0) + 1
    return counts


def build_snapshot(workspace: Workspace, fmt: str, scope: Scope) -> Snapshot:
    """在一个读事务里把导出内容固定下来。

    关键点：**证据行、文献书目、单元页码、集合名称全部走同一个连接**。
    ``export.evidence_view`` 展开出处时用的 ``locator`` / ``get_unit`` /
    ``get_document`` 都复用这条连接（见 :func:`build_snapshot` 传入的 ``conn``），
    不会中途另开连接看到更新的数据。事务结束后内容已复制成脱离数据库的结构，
    摘要与四种格式序列化都只读这份结构。
    """
    if fmt not in FORMATS:
        raise ValidationError(f"不支持的导出格式：{fmt}（可选 {'/'.join(FORMATS)}）")
    with workspace.connect() as conn:
        conn.execute("BEGIN")           # 读事务：固定本次可见的数据版本
        try:
            if scope.collection_id:
                assert_collection_exists(conn, scope.collection_id)
            rows = select_rows(workspace, scope, conn)
            # 同一条连接：书目 / 单元 / 页码 / 定位串与证据行属于同一快照
            repo = evidence_repository(workspace, conn)
            items: list[dict[str, Any]] = []
            for row in rows:
                view = evidence_view(repo, row)
                names = [
                    conn.execute(
                        "SELECT name FROM collections WHERE id = ?", (collection_id,)
                    ).fetchone()["name"]
                    for collection_id in collection_ids_for(row["id"], connection=conn)
                ]
                items.append(_material_item(view, names))
            description = _scope_description(conn, scope)
            documents = _document_summary(items)
            counts = _review_counts(items)
        finally:
            # 只读事务：立刻结束，锁不跨越序列化阶段。
            conn.rollback()
    token = snapshot_token(fmt, scope, description, documents, items)
    return Snapshot(
        format=fmt,
        scope=scope,
        scope_description=description,
        count=len(items),
        review_counts=counts,
        items=tuple(items),
        documents=tuple(documents),
        token=token,
        filename=f"摘录资料_{token[:12]}{EXTENSIONS[fmt]}",
    )


def snapshot_token(fmt: str, scope: Scope, description: str,
                   documents: list[dict[str, Any]], items: list[dict[str, Any]]) -> str:
    """内容摘要：格式 + 范围 + 每条摘录 + 书目页码；**不含生成时间**。"""
    return sha256_text(_canonical({
        "kind": KIND,
        "format": fmt,
        "scope": scope.as_dict(),
        "scope_description": description,
        "documents": documents,
        "items": items,
    }))


# ------------------------------------------------------------------ 渲染


def payload_of(snapshot: Snapshot) -> dict[str, Any]:
    """四种格式共用的结构化内容。"""
    return {
        "kind": KIND,
        "format": snapshot.format,
        "snapshot_token": snapshot.token,
        "scope": snapshot.scope.as_dict(),
        "scope_description": snapshot.scope_description,
        "review_policy": snapshot.scope.review_policy,
        "count": snapshot.count,
        "review_counts": snapshot.review_counts,
        "documents": list(snapshot.documents),
        "items": [dict(item) for item in snapshot.items],
        "material_note": MATERIAL_NOTE,
        "review_note": REVIEW_NOTE,
    }


def render(snapshot: Snapshot) -> bytes:
    payload = payload_of(snapshot)
    if snapshot.format == "json":
        return (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    if snapshot.format == "csv":
        return _render_csv(payload)
    if snapshot.format == "markdown":
        return _render_markdown(payload).encode("utf-8")
    return _render_docx(payload, snapshot)


def _render_csv(payload: dict[str, Any]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    metadata_columns = ("export_scope", "snapshot_token", "material_note", "export_review_note")
    writer.writerow((*MATERIAL_COLUMNS, *metadata_columns))
    metadata = (payload["scope_description"], payload["snapshot_token"], MATERIAL_NOTE, REVIEW_NOTE)
    for item in payload["items"]:
        values = [item.get(column, "") for column in MATERIAL_COLUMNS]
        writer.writerow([csv_safe(_cell(value)) for value in (*values, *metadata)])
    # utf-8-sig：Excel/WPS 直接双击打开不乱码
    return buffer.getvalue().encode("utf-8-sig")


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    return value


def _header_lines(payload: dict[str, Any]) -> list[str]:
    counts = "、".join(
        f"{STATUS_LABELS.get(status, status)} {count}"
        for status, count in payload["review_counts"].items()
    )
    return [
        f"- 导出说明：{md_escape(payload['scope_description'])}",
        f"- 数量：{payload['count']} 条（{md_escape(counts)}）",
        f"- 内容摘要：`{payload['snapshot_token']}`",
    ]


def _render_markdown(payload: dict[str, Any]) -> str:
    lines = [
        "# 摘录资料集",
        "",
        f"- 格式：{FORMAT_LABELS[payload['format']]}",
        *_header_lines(payload),
        "",
        f"> {md_escape(MATERIAL_NOTE)}",
        "",
        f"> {md_escape(REVIEW_NOTE)}",
        "",
    ]
    index = 0
    for document in payload["documents"]:
        lines.append(f"## {md_escape(_document_title(document))}")
        lines.append("")
        lines.append(f"- 文件名：{md_escape(document['filename'])}")
        lines.append(f"- 原件 SHA-256：{md_escape(document['source_sha256'])}")
        lines.append(f"- 本文件含 {document['count']} 条摘录")
        lines.append("")
        for item in payload["items"]:
            if item["doc_id"] != document["doc_id"]:
                continue
            index += 1
            lines.extend(_markdown_item(index, item))
        lines.append("")
    if not payload["items"]:
        lines.extend(["（没有符合条件的摘录）", ""])
    return "\n".join(lines)


def _document_title(document: dict[str, Any]) -> str:
    bits = [document.get("title") or document.get("filename") or document["doc_id"]]
    extra = "，".join(
        str(document.get(field)) for field in ("author", "year", "edition", "volume")
        if document.get(field)
    )
    return f"{bits[0]}（{extra}）" if extra else bits[0]


def _page_label(item: dict[str, Any]) -> str:
    if item["kind"] == "page":
        base = f"PDF 第 {item['file_page']} 页"
    else:
        base = f"Word 文本单元 {item['unit_seq']}"
    if item["printed_page"]:
        base += f"，印刷页码 {item['printed_page']}"
    status = item["printed_page_status"]
    base += {"none": "（未确认）", "candidate": "（候选）",
             "confirmed": "（已确认）"}.get(status, f"（{status}）")
    return base


def _markdown_item(index: int, item: dict[str, Any]) -> list[str]:
    version_label = "校订版" if item["version"] == "corrected" else "识别原文"
    lines = [
        f"### {index}. {md_escape(item['evidence_id'])}｜{md_escape(item['locator'])}",
        "",
        f"- 来源：{md_escape(item['position'])}",
        f"- 页码：{md_escape(_page_label(item))}",
        f"- 单元：{md_escape(item['unit_id'])}｜文字版本：{version_label}"
        f"｜记录时文字哈希：{md_escape(item['unit_text_sha256'] or '（未记录）')}"
        f"｜字符偏移：{item['start_offset']}–{item['end_offset']}",
        f"- 摘录类别：{md_escape(item['category'] or '未分类')}"
        + (f"｜所属集合：{md_escape(item['collections'])}" if item["collections"] else ""),
        f"- 核验状态：{md_escape(item['review_status'])}（{md_escape(item['review_status_label'])}）"
        f"｜核验记录：{md_escape(item['review_note'] or '（无）')}",
    ]
    if item["note"]:
        lines.append(f"- 分析备注：{md_escape(item['note'])}")
    lines.extend(["", "> " + md_escape(item["quote"].replace("\n", "\n> ")), ""])
    if item["warning"]:
        lines.append(f"- ⚠ {md_escape(item['warning'])}")
        lines.append("")
    return lines


def _docx_text(text: str) -> str:
    """可见转义 XML 1.0 禁止的码点，仅影响 Word 呈现，不改保存的证据。"""
    return "".join(
        character if (
            character in "\t\n\r" or 0x20 <= ord(character) <= 0xD7FF
            or 0xE000 <= ord(character) <= 0xFFFD
            or 0x10000 <= ord(character) <= 0x10FFFF
        ) else f"\\u{ord(character):04x}"
        for character in text
    )


def _render_docx(payload: dict[str, Any], snapshot: Snapshot) -> bytes:
    import docx

    document = docx.Document()

    def paragraph(text: str) -> None:
        document.add_paragraph(_docx_text(text))

    def heading(text: str, level: int) -> None:
        document.add_heading(_docx_text(text), level=level)
    heading("摘录资料集", level=0)
    paragraph(f"格式：{FORMAT_LABELS[payload['format']]}")
    for line in _header_lines(payload):
        paragraph(line.replace("`", ""))
    paragraph(MATERIAL_NOTE)
    paragraph(REVIEW_NOTE)

    index = 0
    for entry in payload["documents"]:
        heading(_document_title(entry), level=1)
        paragraph(f"文件名：{entry['filename']}")
        paragraph(f"原件 SHA-256：{entry['source_sha256']}")
        paragraph(f"本文件含 {entry['count']} 条摘录")
        for item in payload["items"]:
            if item["doc_id"] != entry["doc_id"]:
                continue
            index += 1
            version_label = "校订版" if item["version"] == "corrected" else "识别原文"
            heading(f"{index}. {item['evidence_id']}｜{item['locator']}", level=2)
            paragraph(f"引文：{item['quote']}")
            paragraph(
                f"来源：{item['position']}；页码：{_page_label(item)}"
            )
            paragraph(
                f"单元：{item['unit_id']}｜文字版本：{version_label}"
                f"（{item['method']}）｜记录时文字哈希：{item['unit_text_sha256'] or '（未记录）'}"
                f"｜字符偏移：{item['start_offset']}–{item['end_offset']}"
            )
            paragraph(
                f"摘录类别：{item['category'] or '未分类'}"
                + (f"｜所属集合：{item['collections']}" if item["collections"] else "")
            )
            paragraph(
                f"分析备注：{item['note'] or '（无）'}"
            )
            paragraph(
                f"核验状态：{item['review_status']}（{item['review_status_label']}）"
                f"｜核验记录：{item['review_note'] or '（无）'}"
            )
            paragraph(f"内容摘要：{snapshot.token}")
            if item["warning"]:
                paragraph(f"⚠ {item['warning']}")
    if not payload["items"]:
        paragraph("（没有符合条件的摘录）")
    # save() 收的是**二进制流对象**，不是它的字符串化路径：
    # str(BytesIO()) 在 Windows 上是 "<_io.BytesIO object at 0x...>"，不能当路径用。
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


# ------------------------------------------------------------------ 落盘与下载


def exports_dir(workspace: Workspace) -> Path:
    return workspace.artifacts_dir / "exports"


def store_export(workspace: Workspace, snapshot: Snapshot, data: bytes) -> dict[str, Any]:
    """以工作区写事务串行发布；相同快照复用第一次的文件。

    DOCX 的 ZIP 时间戳可以让相同内容产生不同字节，故先读既有登记，
    不能用本次渲染字节否定已发布的同一快照。文件检查、发布与登记都在
    BEGIN IMMEDIATE 内，另一个服务进程也不能覆盖刚登记的文件。
    """
    directory = exports_dir(workspace)
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / snapshot.filename
    with workspace.connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT * FROM exports WHERE filename=?", (snapshot.filename,)).fetchone()
        if row is not None:
            record = _export_record(row)
            if record["snapshot_token"] != snapshot.token or record["format"] != snapshot.format:
                raise EditConflictError("资料文件名称冲突，已中止生成。")
            existing_path = export_path(workspace, record)
            if sha256_file(existing_path) != record["sha256"]:
                raise EditConflictError("已生成的资料文件发生变化，请核对工作区备份。")
            return record

        digest = sha256_bytes(data)
        published = False
        if target.exists():
            if sha256_file(target) != digest:
                raise EditConflictError("未登记的同名资料文件内容不一致，已中止生成。")
        else:
            temporary = directory / f".{snapshot.filename}.{uuid.uuid4().hex[:8]}.tmp"
            try:
                with temporary.open("xb") as handle:
                    handle.write(data)
                temporary.replace(target)
                published = True
            finally:
                temporary.unlink(missing_ok=True)
        record = {
            "id": f"exp_{snapshot.token[:16]}",
            "format": snapshot.format,
            "filename": snapshot.filename,
            "stored_path": target.relative_to(workspace.root).as_posix(),
            "count": snapshot.count,
            "review_counts": snapshot.review_counts,
            "sha256": digest,
            "byte_size": len(data),
            "snapshot_token": snapshot.token,
            "scope": snapshot.scope.as_dict(),
            "created_at": utcnow(),
        }
        try:
            conn.execute(
                "INSERT INTO exports (id, format, filename, stored_path, count, review_counts, "
                "sha256, byte_size, snapshot_token, scope, created_at) "
                "VALUES (:id, :format, :filename, :stored_path, :count, :review_counts, "
                ":sha256, :byte_size, :snapshot_token, :scope, :created_at)",
                {
                    **record,
                    "review_counts": json.dumps(record["review_counts"], ensure_ascii=False),
                    "scope": json.dumps(record["scope"], ensure_ascii=False, sort_keys=True),
                },
            )
            # COMMIT 也可能因读锁超时失败；在写事务仍有效时清理本次文件，
            # 避免释放写锁后误删后继请求成功发布的文件。
            conn.commit()
        except BaseException:
            # 写锁仍在；只清理本次刚发布、尚未登记的文件。
            if published and conn.in_transaction:
                target.unlink(missing_ok=True)
            raise
    return record


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def get_export(workspace: Workspace, export_id: str) -> dict[str, Any]:
    with workspace.connect() as conn:
        row = conn.execute("SELECT * FROM exports WHERE id = ?", (export_id,)).fetchone()
    if row is None:
        raise NotFoundError(f"找不到导出记录：{export_id}")
    return _export_record(row)


def _export_record(row: sqlite3.Row) -> dict[str, Any]:
    record = dict(row)
    record["review_counts"] = json.loads(record["review_counts"] or "{}")
    record["scope"] = json.loads(record["scope"] or "{}")
    return record


def export_path(workspace: Workspace, record: dict[str, Any]) -> Path:
    """按记录解析文件路径，并强制它落在 ``artifacts/exports`` 之内。"""
    directory = exports_dir(workspace).resolve()
    candidate = (workspace.root / str(record["stored_path"])).resolve()
    if not candidate.is_relative_to(directory) or candidate.name != record["filename"]:
        raise NotFoundError("资料文件不在受控目录内，拒绝下载。")
    if not candidate.exists():
        raise NotFoundError(f"资料文件缺失：{record['filename']}")
    return candidate


def export_view(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": record["id"],
        "format": record["format"],
        "filename": record["filename"],
        "count": record["count"],
        "review_counts": record["review_counts"],
        "download_url": f"/api/v1/exports/{record['id']}/download",
        "sha256": record["sha256"],
        "byte_size": record["byte_size"],
        "snapshot_token": record["snapshot_token"],
        "scope": record["scope"],
        "created_at": record["created_at"],
    }


__all__ = [
    "DEFAULT_REVIEW",
    "FORMATS",
    "KIND",
    "MATERIAL_COLUMNS",
    "MATERIAL_NOTE",
    "REVIEW_NOTE",
    "SCOPE_ALL",
    "Snapshot",
    "Scope",
    "build_snapshot",
    "export_path",
    "export_view",
    "exports_dir",
    "get_export",
    "normalize_scope",
    "payload_of",
    "render",
    "snapshot_token",
    "store_export",
]
