"""资料导出：默认只含已接受、四种格式、一致预览与按 ID 下载。

重点覆盖：跨页不截断、显式选择与筛选取交集、空结果、预览后内容变化的 409、
历史引文/哈希/识别方式不被当前版本替换、CSV 公式与中文、Markdown 转义、
DOCX 引文与出处、下载路径越界、重复生成幂等、迁移与服务重启后仍可下载。
"""

from __future__ import annotations

import csv
import io
import json
import zipfile

import pytest
from docx import Document

from litworkbench import collections as core_collections
from litworkbench import evidence_exports as core
from litworkbench.evidence import correct_unit
from litworkbench.store import sha256_text, utcnow


# ------------------------------------------------------------------ 夹具


@pytest.fixture
def seeded(session, run_worker, web_workspace, web_fixtures):
    response = session.upload([("files", ("县志.pdf", web_fixtures["text_pdf"].read_bytes(),
                                         "application/pdf"))])
    assert response.status_code == 202
    run_worker()
    doc_id = web_workspace.list_documents()[0]["id"]
    web_workspace.update_document(doc_id, title="合成县志", author="验收编者", year="2026",
                                  edition="验收版", volume="测试卷")
    units = web_workspace.units_for_doc(doc_id)
    return {"doc_id": doc_id, "units": units, "workspace": web_workspace}


def make_evidence(workspace, unit, index: int, *, status: str = "pending",
                  category: str = "", note: str = "", review_note: str = "",
                  method: str = "text") -> str:
    text = unit["text"]
    quote = text[0:4]
    evidence_id = f"ev_x{index:04d}"
    snapshot = {
        "version": "original",
        "unit_text_sha256": sha256_text(text),
        "unit_text_length": len(text),
        "method": method,
        "captured_at": utcnow(),
    }
    with workspace.connect() as conn:
        conn.execute(
            "INSERT INTO evidence (id, unit_id, doc_id, task_id, quote, start_offset, "
            "end_offset, version, version_snapshot, category, note, review_status, "
            "review_note, created_at, updated_at) "
            "VALUES (?, ?, ?, NULL, ?, 0, 4, 'original', ?, ?, ?, ?, ?, ?, ?)",
            (evidence_id, unit["id"], unit["doc_id"], quote,
             json.dumps(snapshot, ensure_ascii=False), category, note, status,
             review_note, utcnow(), utcnow()),
        )
    return evidence_id


@pytest.fixture
def mixed(session, seeded):
    """65 条摘录，状态混合（accepted/pending/rejected/stale），分布在 3 个页单元上。"""
    workspace = seeded["workspace"]
    units = seeded["units"]
    statuses = ["accepted", "accepted", "accepted", "pending", "rejected", "stale"]
    ids: list[str] = []
    for index in range(65):
        unit = units[index % len(units)]
        ids.append(make_evidence(workspace, unit, index, status=statuses[index % len(statuses)],
                                 category="文物" if index % 2 == 0 else "古迹",
                                 note=f"备注 {index}", review_note="人工看过" if index % 3 == 0 else ""))
    return {**seeded, "ids": ids, "statuses": statuses}


def preview(session, fmt="json", **body):
    return session.post("/api/v1/exports/preview", json={"format": fmt, **body})


def generate(session, token, fmt="json", **body):
    return session.post("/api/v1/exports",
                        json={"format": fmt, "snapshot_token": token, **body})


def make_export(session, fmt="json", **body):
    """预览 → 生成（同一种格式：摘要覆盖格式，跨格式预览必然冲突）。"""
    token = preview(session, fmt=fmt, **body).json()["snapshot_token"]
    response = generate(session, token, fmt=fmt, **body)
    assert response.status_code == 201, response.text
    return response.json()


def body_token(row: dict[str, str], rows: list[dict[str, str]]) -> str:
    """CSV 每行都带同一个内容摘要列：整份文件只对应一个快照。"""
    tokens = {item.get("snapshot_token", "") for item in rows}
    assert tokens == {row["snapshot_token"]}
    return row["snapshot_token"]


# ------------------------------------------------------------------ 范围口径


def test_default_scope_is_accepted_only(session, mixed):
    response = preview(session, filters={"review_status": None})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["count"] == 33                    # 65 条里 accepted 的数量
    assert body["review_counts"]["accepted"] == 33
    assert body["review_counts"]["pending"] == 0
    assert body["review_policy"] == "accepted"
    assert "仅已接受（默认）" in body["scope_description"]

    # 未指定 review_status 与显式 accepted 完全等价
    assert preview(session).json()["count"] == body["count"]
    assert preview(session, filters={"review_status": "accepted"}).json()["snapshot_token"] \
        == body["snapshot_token"]


def test_normalize_scope_defaults_to_accepted():
    """范围归一化：不给就等于 accepted，只有显式 all/具体状态才绕过。"""
    assert core.normalize_scope().review_status == "accepted"
    assert core.normalize_scope(review_status=None).review_status == "accepted"
    assert core.normalize_scope(review_status="  ").review_status == "accepted"
    assert core.normalize_scope(review_status="all").review_policy == "all"
    assert core.normalize_scope(review_status="stale").review_status == "stale"
    with pytest.raises(Exception):
        core.normalize_scope(review_status="good")
    with pytest.raises(Exception):
        core.normalize_scope(fmt="pdf")


def test_explicit_all_includes_every_status_with_warnings(session, mixed):
    body = preview(session, filters={"review_status": "all"}).json()
    assert body["count"] == 65
    assert body["review_counts"] == {"pending": 11, "accepted": 33, "rejected": 11,
                                     "stale": 10}

    record = generate(session, body["snapshot_token"],
                      filters={"review_status": "all"}).json()
    payload = json.loads(download_bytes(session, record["id"]))
    by_status = {item["evidence_id"]: item for item in payload["items"]}
    assert payload["review_counts"]["rejected"] == 11
    rejected = [item for item in payload["items"] if item["review_status"] == "rejected"]
    assert rejected and all("不应引用" in item["warning"] for item in rejected)
    assert all(item["usable"] is False for item in rejected)
    stale = [item for item in payload["items"] if item["review_status"] == "stale"]
    assert all("复核" in item["warning"] for item in stale)
    assert by_status[mixed["ids"][0]]["review_status"] == "accepted"


def test_filters_and_explicit_selection_intersect(session, mixed):
    accepted_ids = [mixed["ids"][index] for index in range(65) if index % 6 in (0, 1, 2)]
    pending_ids = [mixed["ids"][index] for index in range(65) if index % 6 == 3]
    rejected_ids = [mixed["ids"][index] for index in range(65) if index % 6 == 4]

    # 默认 accepted ∩ 显式选择：pending 被挡掉
    body = preview(session, evidence_ids=accepted_ids + pending_ids).json()
    assert body["count"] == len(accepted_ids)
    # 显式 all 时交集里含被拒绝的条目
    every = preview(session, filters={"review_status": "all"},
                    evidence_ids=rejected_ids + pending_ids).json()
    assert every["count"] == len(rejected_ids) + len(pending_ids)
    # 分类条件再取交集（pending 全部落在「古迹」分类上）
    both = preview(session, filters={"category": "古迹", "review_status": "all"},
                   evidence_ids=pending_ids).json()
    assert both["count"] == len(pending_ids)
    assert "分类：古迹" in both["scope_description"]
    other = preview(session, filters={"category": "文物", "review_status": "all"},
                    evidence_ids=pending_ids).json()
    assert other["count"] == 0


def test_collection_filter_scopes_export(session, mixed):
    workspace = mixed["workspace"]
    collection = core_collections.create_collection(workspace, "文物建筑资料")
    ids = mixed["ids"][:20]
    core_collections.apply_batch(workspace, core_collections.plan_batch(
        evidence_ids=ids, add_to_collection_id=collection["id"]))
    body = preview(session, filters={"collection_id": collection["id"],
                                     "review_status": "all"}).json()
    assert body["count"] == len(ids)
    assert "文物建筑资料" in body["scope_description"]


def test_unknown_collection_is_rejected(session, mixed):
    assert preview(session, filters={"collection_id": "col_missing"}).status_code == 404
    assert preview(session, filters={"doc_id": "doc_missing"}).status_code == 200  # 0 条


def test_unknown_explicit_id_is_rejected_not_ignored(session, mixed):
    response = preview(session, evidence_ids=[mixed["ids"][0], "ev_not_here"])
    assert response.status_code == 422
    assert "找不到这些摘录" in response.json()["error"]["message"]
    generated = generate(session, "x" * 64, evidence_ids=[mixed["ids"][0], "ev_not_here"])
    assert generated.status_code == 422


def test_empty_scope_previews_zero_but_refuses_to_generate(session, mixed):
    body = preview(session, filters={"review_status": "accepted", "category": "不存在的分类"})
    assert body.status_code == 200 and body.json()["count"] == 0
    refused = generate(session, body.json()["snapshot_token"],
                       filters={"review_status": "accepted", "category": "不存在的分类"})
    assert refused.status_code == 422
    assert "没有可导出的摘录" in refused.json()["error"]["message"]


@pytest.mark.parametrize("payload,fragment", [
    ({"evidence_ids": []}, "0 条"),
    ({"evidence_ids": ["ev_x0000", ""]}, "非空"),
])
def test_explicit_empty_selection_never_falls_back_to_all(session, mixed, payload, fragment):
    response = preview(session, **payload)
    assert response.status_code == 422
    assert fragment in response.json()["error"]["message"]
    # 对照：完全不传 evidence_ids 才是「当前筛选的全部结果」
    assert preview(session).json()["count"] > 0


def test_export_rejects_more_than_500_raw_ids(session, mixed):
    ids = [mixed["ids"][0]] * 501
    response = preview(session, evidence_ids=ids)
    assert response.status_code == 422
    assert "500" in response.json()["error"]["message"]


def test_export_deduplicates_ids_by_first_occurrence(session, mixed):
    ids = [mixed["ids"][index] for index in (3, 0, 4)]
    body = preview(session, filters={"review_status": "all"},
                   evidence_ids=[ids[2], ids[0], ids[2], ids[1], ids[0]]).json()
    assert body["count"] == 3
    assert body["scope"]["evidence_ids"] == [ids[2], ids[0], ids[1]]
    record = make_export(session, filters={"review_status": "all"}, evidence_ids=ids)
    payload = json.loads(download_bytes(session, record["id"]))
    assert payload["count"] == 3
    assert len(payload["items"]) == 3


def test_unknown_request_fields_are_rejected(session, mixed):
    response = session.post("/api/v1/exports/preview", json={
        "format": "json", "path": "C:/任意/路径.csv",
    })
    assert response.status_code == 422
    assert generate(session, "0" * 64, output="x").status_code == 422


# ------------------------------------------------------------------ 大范围与一致性


def test_full_filtered_export_is_not_truncated_at_30(session, mixed):
    """列表接口一页 30 条；筛选式导出必须覆盖全部 65 条。"""
    body = preview(session, filters={"review_status": "all"}).json()
    assert body["count"] == 65
    record = generate(session, body["snapshot_token"], filters={"review_status": "all"}).json()
    assert record["count"] == 65
    payload = json.loads(download_bytes(session, record["id"]))
    assert len(payload["items"]) == 65
    assert len({item["evidence_id"] for item in payload["items"]}) == 65
    # 列表接口本身仍然是分页的
    assert session.get("/api/v1/evidence?page_size=30").json()["total"] == 65


def test_preview_then_content_change_returns_409(session, mixed):
    workspace = mixed["workspace"]
    body = preview(session).json()
    workspace.get_document(mixed["doc_id"])
    # 审核状态变化
    with workspace.connect() as conn:
        conn.execute("UPDATE evidence SET review_status = 'rejected', updated_at = ? "
                     "WHERE id = ?", (utcnow(), mixed["ids"][0]))
    conflict = generate(session, body["snapshot_token"])
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "EDIT_CONFLICT"

    # 重新预览后恢复
    again = preview(session).json()
    assert again["snapshot_token"] != body["snapshot_token"]
    assert again["count"] == 32
    assert generate(session, again["snapshot_token"]).status_code == 201


def test_correction_between_preview_and_generate_returns_409(session, mixed):
    workspace = mixed["workspace"]
    unit = mixed["units"][0]
    body = preview(session, filters={"review_status": "all"}).json()
    correct_unit(workspace, unit["id"], unit["text"] + "（校订补字）", "补字")
    assert generate(session, body["snapshot_token"],
                    filters={"review_status": "all"}).status_code == 409
    refreshed = preview(session, filters={"review_status": "all"}).json()
    assert refreshed["review_counts"]["stale"] >= 1
    assert refreshed["snapshot_token"] != body["snapshot_token"]
    assert generate(session, refreshed["snapshot_token"],
                    filters={"review_status": "all"}).status_code == 201
    # 默认口径下这条摘录已经不在范围里：数量减少而不是状态被悄悄改写
    assert preview(session).json()["count"] < body["count"]


def test_bibliography_change_between_preview_and_generate_returns_409(session, mixed):
    workspace = mixed["workspace"]
    body = preview(session).json()
    workspace.update_document(mixed["doc_id"], title="改名之后的书名")
    assert generate(session, body["snapshot_token"]).status_code == 409
    assert generate(session, preview(session).json()["snapshot_token"]).status_code == 201


def test_generated_file_is_detached_snapshot(session, mixed):
    """算完摘要之后改书目与校订原文，序列化仍必须用同一份快照。

    摘要计算、四种格式序列化都只读 :class:`Snapshot`，不再查数据库，
    因此不会出现「旧书目 + 新文字」这种混排。
    """
    workspace = mixed["workspace"]
    scope = core.normalize_scope(review_status="all")
    snapshot = core.build_snapshot(workspace, "json", scope)
    old_title = snapshot.documents[0]["title"]
    old_quote = snapshot.items[0]["quote"]
    old_hash = snapshot.items[0]["unit_text_sha256"]

    # 摘要算完之后才改动数据
    workspace.update_document(mixed["doc_id"], title="改动后的书名")
    correct_unit(workspace, mixed["units"][0]["id"],
                 mixed["units"][0]["text"] + "（校订）", "校订")

    # 摘要不变（快照没有被改写）
    assert snapshot.documents[0]["title"] == old_title
    assert snapshot.token == core.snapshot_token(
        snapshot.format, scope, snapshot.scope_description,
        list(snapshot.documents), list(snapshot.items))
    # 四种格式都渲染自这份快照
    payload = json.loads(core.render(snapshot))
    assert payload["documents"][0]["title"] == old_title
    assert payload["items"][0]["quote"] == old_quote
    assert payload["items"][0]["unit_text_sha256"] == old_hash
    markdown = core.render(snapshot).decode("utf-8")
    assert "改动后的书名" not in markdown
    assert old_title in markdown


def test_old_artifact_keeps_content_after_bibliography_change(session, mixed):
    workspace = mixed["workspace"]
    body = preview(session, filters={"review_status": "all"}).json()
    record = generate(session, body["snapshot_token"],
                      filters={"review_status": "all"}).json()
    workspace.update_document(mixed["doc_id"], title="后来改的书名")
    payload = json.loads(download_bytes(session, record["id"]))
    assert payload["documents"][0]["title"] == "合成县志"


def test_repeated_generate_and_download_are_idempotent(session, mixed):
    body = preview(session, filters={"review_status": "all"}).json()
    first = generate(session, body["snapshot_token"], filters={"review_status": "all"}).json()
    second = generate(session, body["snapshot_token"], filters={"review_status": "all"})
    assert second.status_code == 201, second.text
    assert second.json()["id"] == first["id"]
    assert second.json()["filename"] == first["filename"]
    assert second.json()["created_at"] == first["created_at"]
    assert download_bytes(session, first["id"]) == download_bytes(session, second.json()["id"])
    files = sorted(p.name for p in core.exports_dir(mixed["workspace"]).iterdir())
    assert files == [first["filename"]]


def test_historical_quote_hash_and_method_are_not_replaced(session, mixed):
    """引文、偏移、文字哈希与识别方式都取自记录时的版本快照。"""
    workspace = mixed["workspace"]
    unit = mixed["units"][0]
    evidence_id = make_evidence(workspace, unit, 900, status="accepted", method="text")
    recorded_hash = sha256_text(unit["text"])
    quote = unit["text"][0:4]

    # 校订 + 改识别方式：当前单元的 method 变成 ocr
    unit_row = workspace.get_unit(unit["id"])
    unit_row["method"] = "ocr"
    workspace.save_unit(unit_row)
    correct_unit(workspace, unit["id"], unit["text"] + "（后来的校订）", "补字")

    body = preview(session, filters={"review_status": "all"}).json()
    record = generate(session, body["snapshot_token"], filters={"review_status": "all"}).json()
    payload = json.loads(download_bytes(session, record["id"]))
    item = next(row for row in payload["items"] if row["evidence_id"] == evidence_id)
    assert item["quote"] == quote                  # 历史引文未被当前全文替换
    assert item["start_offset"] == 0 and item["end_offset"] == 4
    assert item["version"] == "original"
    assert item["unit_text_sha256"] == recorded_hash
    assert item["method"] == "text"                # 记录时的识别方式，不是当前的 ocr
    assert item["review_status"] == "stale"        # 当前核验状态取自一致快照
    assert "复核" in item["warning"]


def test_legacy_snapshot_without_method_falls_back(session, mixed):
    workspace = mixed["workspace"]
    unit = mixed["units"][1]
    evidence_id = f"ev_legacy_{unit['id']}"
    with workspace.connect() as conn:
        conn.execute(
            "INSERT INTO evidence (id, unit_id, doc_id, task_id, quote, start_offset, "
            "end_offset, version, version_snapshot, category, note, review_status, "
            "review_note, created_at, updated_at) "
            "VALUES (?, ?, ?, NULL, '旧记录', 0, 3, 'original', '{}', '', '', 'accepted', "
            "'', ?, ?)", (evidence_id, unit["id"], unit["doc_id"], utcnow(), utcnow()))
    scope = core.normalize_scope(review_status="all")
    snapshot = core.build_snapshot(workspace, "json", scope)
    item = next(row for row in snapshot.items if row["evidence_id"] == evidence_id)
    assert item["method"] == "text"                # 极老记录才退回单元当前值
    assert item["unit_text_sha256"] == ""


# ------------------------------------------------------------------ 四种格式


def download_bytes(session, export_id: str) -> bytes:
    response = session.get(f"/api/v1/exports/{export_id}/download")
    assert response.status_code == 200, response.text
    return response.content


def test_csv_has_bom_provenance_and_formula_escape(session, mixed):
    """CSV 的第一行就是真表头（标准 csv.DictReader 可直接读）。

    资料元信息（范围、核验口径、内容摘要）作为普通列出现在每一行里，
    不使用 ``#`` 注释前言——那会让多数表格工具把表头认错。
    """
    workspace = mixed["workspace"]
    hostile = make_evidence(workspace, mixed["units"][0], 950, status="accepted",
                            category="=cmd|'/c calc'!A1", note="+1+1")
    record = make_export(session, fmt="csv", filters={"review_status": "all"})
    raw = download_bytes(session, record["id"])
    assert raw.startswith(b"\xef\xbb\xbf")                 # UTF-8 BOM

    text = raw.decode("utf-8-sig")
    first_line = text.splitlines()[0]
    assert first_line.startswith("evidence_id,")            # 没有注释前言
    rows = list(csv.DictReader(io.StringIO(text)))
    assert len(rows) == 66                                 # 夹具 65 条 + 本用例新增的 1 条
    assert "printed_page_status" in first_line
    assert "unit_text_sha256" in first_line

    record_row = next(row for row in rows if row["evidence_id"] == hostile)
    assert record_row["category"].startswith("'=")          # 公式被中和
    assert record_row["note"].startswith("'+")             # 含 + 的文本同样被中和
    assert record_row["source_sha256"] == workspace.get_document(
        mixed["doc_id"])["sha256"]
    assert "合成县志" in text and "验收编者" in text
    # 元信息作为普通列随每一行出现（与其它列一样，不使用 # 注释前言）
    assert record_row["review_status"] == "accepted"
    assert record_row["export_scope"]
    assert record_row["snapshot_token"] == body_token(record_row, rows)
    assert "摘录资料集" in record_row["material_note"]
    assert record_row["export_review_note"]


def test_markdown_escapes_and_groups_by_document(session, mixed):
    workspace = mixed["workspace"]
    with workspace.connect() as conn:
        conn.execute("UPDATE units SET printed_page = '775' WHERE doc_id = ?",
                     (mixed["doc_id"],))
    hostile = make_evidence(workspace, mixed["units"][0], 960, status="accepted",
                            note="含 | 竖线 与 *星号*")
    text = download_bytes(session, make_export(
        session, fmt="markdown", filters={"review_status": "all"})["id"]).decode("utf-8")
    assert "\\| 竖线" in text and "\\*星号\\*" in text
    assert "## 合成县志" in text
    assert hostile in text
    assert "原件 SHA-256" in text
    assert "印刷页码 775" in text and "未确认" in text
    assert "内容摘要" in text
    assert "摘录资料集" in text and "不包含研究结论" in text
    assert "- 范围：范围：" not in text          # 标题前缀不重复


def test_docx_contains_quote_and_provenance(session, mixed):
    with mixed["workspace"].connect() as conn:
        conn.execute("UPDATE units SET printed_page = '775' WHERE doc_id = ?",
                     (mixed["doc_id"],))
    raw = download_bytes(session, make_export(
        session, fmt="docx", filters={"review_status": "all"})["id"])
    assert raw[:2] == b"PK"                             # 真正的 zip/OOXML 容器
    assert zipfile.is_zipfile(io.BytesIO(raw))
    text = "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
    assert "摘录资料集" in text
    assert "合成县志" in text and "验收编者" in text and "验收版" in text and "测试卷" in text
    assert mixed["workspace"].get_document(mixed["doc_id"])["sha256"] in text
    assert "印刷页码 775" in text and "未确认" in text
    assert "记录时文字哈希" in text and "核验状态" in text and "字符偏移" in text
    assert mixed["units"][0]["text"][0:4] in text


def test_printed_page_status_present_in_every_format(session, mixed):
    unit_id = mixed["units"][0]["id"]
    with mixed["workspace"].connect() as conn:
        conn.execute("UPDATE units SET printed_page = '775', printed_page_status = "
                     "'confirmed' WHERE id = ?", (unit_id,))

    csv_text = download_bytes(session, make_export(
        session, fmt="csv", filters={"review_status": "all"})["id"]).decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(csv_text)))
    row = next(item for item in rows if item["unit_id"] == unit_id)
    assert row["printed_page"] == "775" and row["printed_page_status"] == "confirmed"

    markdown = download_bytes(session, make_export(
        session, fmt="markdown", filters={"review_status": "all"})["id"]).decode("utf-8")
    assert "印刷页码 775" in markdown and "已确认" in markdown

    payload = json.loads(download_bytes(session, make_export(
        session, fmt="json", filters={"review_status": "all"})["id"]))
    item = next(row for row in payload["items"] if row["unit_id"] == unit_id)
    assert item["printed_page"] == "775" and item["printed_page_status"] == "confirmed"

    docx_text = "\n".join(p.text for p in Document(io.BytesIO(download_bytes(
        session,
        make_export(session, fmt="docx", filters={"review_status": "all"})["id"],
    ))).paragraphs)
    assert "印刷页码 775" in docx_text and "已确认" in docx_text


def test_no_local_source_path_leaks_into_exports(session, mixed):
    workspace = mixed["workspace"]
    absolute = str(workspace.path_of(workspace.get_document(mixed["doc_id"])))
    json_text = download_bytes(session, make_export(
        session, fmt="json", filters={"review_status": "all"})["id"]).decode("utf-8")
    markdown = download_bytes(session, make_export(
        session, fmt="markdown", filters={"review_status": "all"})["id"]).decode("utf-8")
    for text in (json_text, markdown):
        assert absolute not in text
        assert str(workspace.root) not in text
    assert "source_path" not in json_text
    assert json.loads(json_text)["documents"][0]["source_sha256"]


def test_export_payload_declares_material_set_not_report(session, mixed):
    body = preview(session).json()
    assert "不包含研究结论" in body["material_note"]
    payload = json.loads(download_bytes(session, make_export(session)["id"]))
    assert payload["kind"] == core.KIND
    assert set(payload) >= {"kind", "format", "snapshot_token", "scope", "count",
                            "review_counts", "documents", "items", "material_note",
                            "review_note"}
    assert "findings" not in payload and "conclusions" not in payload
    for item in payload["items"]:
        assert "source_path" not in item
        assert set(item) >= set(core.MATERIAL_COLUMNS)


# ------------------------------------------------------------------ 下载安全


def test_download_requires_session_and_rejects_cross_origin(app, mixed):
    from fastapi.testclient import TestClient

    body = None
    with TestClient(app, base_url="http://127.0.0.1:8765") as fresh:
        token = fresh.get("/api/v1/workspace").json()["csrf_token"]
        headers = {"X-CSRF-Token": token}
        body = fresh.post("/api/v1/exports/preview", json={"format": "json"},
                          headers=headers).json()
        record = fresh.post("/api/v1/exports",
                            json={"format": "json", "snapshot_token": body["snapshot_token"]},
                            headers=headers).json()
        denied = fresh.post("/api/v1/exports/preview", json={"format": "json"},
                            headers={"X-CSRF-Token": token,
                                     "Origin": "http://evil.example"})
        assert denied.status_code == 403
        no_csrf = fresh.post("/api/v1/exports/preview", json={"format": "json"})
        assert no_csrf.status_code == 403
        assert fresh.get(f"/api/v1/exports/{record['id']}/download").status_code == 200
    with TestClient(app, base_url="http://127.0.0.1:8765") as stranger:
        assert stranger.get(f"/api/v1/exports/{record['id']}/download").status_code == 401


def test_download_headers_and_unknown_ids(session, mixed):
    from urllib.parse import unquote

    record = make_export(session)
    response = session.get(f"/api/v1/exports/{record['id']}/download")
    assert response.headers["content-type"].startswith("application/json")
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment")
    assert "摘录资料_" in unquote(disposition)      # RFC 5987 编码后的中文名
    assert response.headers["x-export-sha256"] == record["sha256"]
    detail = session.get(f"/api/v1/exports/{record['id']}").json()
    assert detail["download_url"] == f"/api/v1/exports/{record['id']}/download"
    assert session.get("/api/v1/exports/exp_missing/download").status_code == 404
    assert session.get("/api/v1/exports/exp_missing").status_code == 404


def test_download_rejects_paths_outside_exports_dir(session, mixed):
    record = make_export(session)
    outside = mixed["workspace"].root / "sources" / "leak.csv"
    outside.write_text("不该被下载", encoding="utf-8")
    with mixed["workspace"].connect() as conn:
        conn.execute("UPDATE exports SET stored_path = ? WHERE id = ?",
                     (outside.relative_to(mixed["workspace"].root).as_posix(), record["id"]))
    response = session.get(f"/api/v1/exports/{record['id']}/download")
    assert response.status_code == 404
    assert "受控目录" in response.json()["error"]["message"]

    # 绝对路径同样不行
    with mixed["workspace"].connect() as conn:
        conn.execute("UPDATE exports SET stored_path = ? WHERE id = ?",
                     (str(outside), record["id"]))
    assert session.get(f"/api/v1/exports/{record['id']}/download").status_code == 404


def test_download_rejects_filename_mismatch(session, mixed):
    record = make_export(session)
    with mixed["workspace"].connect() as conn:
        conn.execute("UPDATE exports SET filename = ? WHERE id = ?",
                     ("../../escape.csv", record["id"]))
    assert session.get(f"/api/v1/exports/{record['id']}/download").status_code == 404


def test_export_survives_service_restart(session, mixed):
    record = make_export(session)
    before = download_bytes(session, record["id"])

    from fastapi.testclient import TestClient

    from litworkbench.web.app import create_app

    restarted = create_app(mixed["workspace"].root, start_worker=False)
    with TestClient(restarted, base_url="http://127.0.0.1:8765") as fresh:
        fresh.get("/api/v1/workspace")
        response = fresh.get(f"/api/v1/exports/{record['id']}/download")
        assert response.status_code == 200
        assert response.content == before
        assert fresh.get("/api/v1/collections").status_code == 200
        assert fresh.get("/api/v1/evidence?page_size=5").json()["total"] == 65


def test_migrated_workspace_can_export_without_new_tables(seeded, tmp_path):
    """从旧库（没有 exports 表）升级后即可导出，旧摘录照样能用。"""
    from litworkbench import migrations
    from litworkbench.store import Workspace

    root = tmp_path / "迁移 导出 工作区"
    workspace = Workspace(root)
    workspace.initialize()
    with workspace.connect() as conn:
        conn.execute("DROP TABLE exports")
        conn.execute("UPDATE meta SET value = '5' WHERE key = 'schema_version'")
    migrations.ensure_schema(workspace)
    scope = core.normalize_scope()
    snapshot = core.build_snapshot(workspace, "csv", scope)
    assert snapshot.count == 0
    record = core.store_export(workspace, snapshot, core.render(snapshot))
    assert core.export_path(workspace, core.get_export(workspace, record["id"])).exists()